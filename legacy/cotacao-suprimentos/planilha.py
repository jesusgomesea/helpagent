# -*- coding: utf-8 -*-
"""Exporta o resultado da cotacao para xlsx, no mesmo formato acordado no piloto."""
import datetime
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

AZUL = "1F3864"
AZUL_2 = "2E5C8A"
CINZA = "F2F2F2"
VERDE = "E2EFDA"
AMARELO = "FFF2CC"
LARANJA = "FCE4D6"
FT = "Arial"


def _titulo(ws, cel, texto, cor=AZUL):
    ws[cel] = texto
    ws[cel].font = Font(name=FT, size=12, bold=True, color="FFFFFF")
    ws[cel].fill = PatternFill("solid", fgColor=cor)


def _cabecalho(ws, linha, headers, larguras):
    for i, (h, w) in enumerate(zip(headers, larguras), 1):
        c = ws.cell(row=linha, column=i, value=h)
        c.font = Font(name=FT, size=9, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[linha].height = 28


def gerar(termo, resultado):
    r = resultado
    res = r["resumo"]
    crit = r.get("criterios", {})
    hoje = datetime.date.today()
    wb = Workbook()

    # ---------------------------------------------------------------- resumo
    rs = wb.active
    rs.title = "Resumo"
    rs.sheet_view.showGridLines = False
    _titulo(rs, "A1", "COTACAO  -  %s" % termo.upper())
    rs.merge_cells("A1:H1")
    rs.row_dimensions[1].height = 24

    meta = [("Termo pesquisado", termo), ("Fonte", "Mercado Livre Brasil"),
            ("Data da coleta", hoje.strftime("%d/%m/%Y")),
            ("Anuncios analisados", res["analisados"]),
            ("Elegiveis apos eliminatorios", res["elegiveis"]),
            ("Descartados", res["descartados"]),
            ("Menor preco elegivel", res["menor_preco"]),
            ("Preco medio elegivel", res["preco_medio"])]
    for i, (k, v) in enumerate(meta, start=3):
        rs["A%d" % i] = k
        rs["A%d" % i].font = Font(name=FT, size=10, bold=True)
        rs["C%d" % i] = v
        rs["C%d" % i].font = Font(name=FT, size=10)
        if "preco" in k.lower():
            rs["C%d" % i].number_format = '"R$" #,##0.00'

    linha = 13
    for g in r["grupos"]:
        _titulo(rs, "A%d" % linha,
                "TOP 3  -  %s   (%d elegiveis)" % (g["rotulo"], g["qtd"]), AZUL_2)
        rs.merge_cells("A%d:H%d" % (linha, linha))
        linha += 1
        _cabecalho(rs, linha, ["Posicao", "Produto", "Preco", "Score", "Nota",
                               "Vendidos", "Entrega", "Link"],
                   [9, 52, 13, 9, 7, 10, 14, 15])
        linha += 1
        for pos, a in enumerate(g["top3"], 1):
            rs.cell(linha, 1, "#%d" % pos).font = Font(name=FT, size=11, bold=True)
            rs.cell(linha, 2, a["titulo"])
            rs.cell(linha, 3, a["preco"]).number_format = '"R$" #,##0.00'
            rs.cell(linha, 4, a["score"] / 100).number_format = "0.0%"
            rs.cell(linha, 5, a["nota"]).number_format = "0.0"
            rs.cell(linha, 6, a["vendidos"])
            rs.cell(linha, 7, "FULL" if a["full"] else
                    ("Frete gratis" if a["frete_gratis"] else "Envio comum"))
            c = rs.cell(linha, 8, ("abrir anuncio (ad)" if a.get("patrocinado")
                                   else "abrir anuncio") if a["url"] else "sem link")
            if a["url"]:
                c.hyperlink = a["url"]
                c.font = Font(name=FT, size=10, color="0563C1", underline="single")
            fill = VERDE if pos == 1 else "FFFFFF"
            for col in range(1, 9):
                cl = rs.cell(linha, col)
                if not (col == 8 and a["url"]) and col != 1:
                    cl.font = Font(name=FT, size=10, bold=(pos == 1))
                cl.fill = PatternFill("solid", fgColor=fill)
                cl.border = Border(bottom=Side("thin", color="D9D9D9"))
            linha += 1
        rs.cell(linha, 2, "Menor preco do grupo").font = Font(name=FT, size=9, italic=True)
        rs.cell(linha, 3, g["menor_preco"]).number_format = '"R$" #,##0.00'
        rs.cell(linha, 3).font = Font(name=FT, size=9, italic=True)
        rs.cell(linha, 5, "Preco medio").font = Font(name=FT, size=9, italic=True)
        rs.cell(linha, 6, g["preco_medio"]).number_format = '"R$" #,##0.00'
        rs.cell(linha, 6).font = Font(name=FT, size=9, italic=True)
        linha += 3

    rs["A%d" % linha] = ("Precos coletados em %s. O Mercado Livre exibe precos diferentes "
                         "conforme a conta logada e o nivel Meli+; confirme no anuncio antes "
                         "de fechar a compra." % hoje.strftime("%d/%m/%Y"))
    rs["A%d" % linha].font = Font(name=FT, size=9, italic=True, color="C00000")
    rs.merge_cells("A%d:H%d" % (linha, linha))
    for col, w in zip("ABCDEFGH", [10, 52, 13, 10, 8, 11, 15, 16]):
        rs.column_dimensions[col].width = w

    # -------------------------------------------------------------- criterios
    cr = wb.create_sheet("Criterios")
    cr.sheet_view.showGridLines = False
    _titulo(cr, "A1", "CRITERIOS USADOS NESTA COTACAO")
    cr.merge_cells("A1:C1")
    cr["A2"] = ("Estes sao os valores que o usuario aplicou na pagina no momento da "
                "cotacao. Mudar o criterio muda o ranking - refaca a cotacao.")
    cr["A2"].font = Font(name=FT, size=9, italic=True)

    pe = res["pesos_efetivos"]
    blocos = [
        ("PESOS EFETIVOS", [("Preco", "%d%%" % pe["peso_preco"]),
                            ("Entrega", "%d%%" % pe["peso_entrega"]),
                            ("Fornecedor", "%d%%" % pe["peso_fornecedor"]),
                            ("Marca", "%d%%" % pe["peso_marca"])]),
        ("ELIMINATORIOS", [("Nota minima do vendedor", crit.get("nota_minima")),
                           ("Minimo de unidades vendidas", crit.get("vendidos_minimo")),
                           ("Aceita recondicionado", "sim" if crit.get("aceita_recondicionado") else "nao"),
                           ("Exige envio FULL", "sim" if crit.get("exigir_full") else "nao"),
                           ("Origem do envio", {"qualquer": "qualquer origem",
                                                "nacional": "so nacional",
                                                "internacional": "so internacional"}
                            .get(crit.get("origem"), crit.get("origem"))),
                           ("Teto de preco", crit.get("preco_max") or "sem teto"),
                           ("Titulo deve conter", crit.get("deve_conter") or "-"),
                           ("Titulo nao pode conter", crit.get("nao_pode_conter") or "-")]),
        ("PONTUACAO DE ENTREGA", [("Enviado pelo FULL", "%s pts" % crit.get("pontos_full")),
                                  ("Frete gratis", "%s pts" % crit.get("pontos_frete_gratis")),
                                  ("Sem frete gratis", "%s pts" % crit.get("pontos_sem_frete"))]),
        ("OUTROS", [("Marcas preferidas", crit.get("marcas_preferidas") or "lista padrao de TI"),
                    ("Segmentos", crit.get("segmentos") or "sem segmentacao")]),
    ]
    i = 4
    for nome, itens in blocos:
        cr["A%d" % i] = nome
        cr["A%d" % i].font = Font(name=FT, size=10, bold=True)
        i += 1
        for k, v in itens:
            cr["A%d" % i] = k
            cr["A%d" % i].font = Font(name=FT, size=10)
            cr["B%d" % i] = v
            cr["B%d" % i].font = Font(name=FT, size=10, bold=True, color="0000FF")
            cr["B%d" % i].fill = PatternFill("solid", fgColor=AMARELO)
            i += 1
        i += 1
    cr.column_dimensions["A"].width = 34
    cr.column_dimensions["B"].width = 40

    # ---------------------------------------------------------------- analise
    an = wb.create_sheet("Analise")
    an.sheet_view.showGridLines = False
    heads = ["Produto", "Segmento", "Preco", "Preco de", "Nota", "Vendidos", "FULL",
             "Origem", "Loja oficial", "Frete gratis", "Marca", "Vendedor", "Link",
             "Pts preco", "Pts entrega", "Pts fornecedor", "Pts marca", "SCORE"]
    _cabecalho(an, 1, heads,
               [52, 14, 12, 12, 7, 10, 7, 14, 11, 11, 12, 20, 14, 10, 11, 12, 10, 10])
    for i, a in enumerate(r["elegiveis"], start=2):
        p = a["parciais"]
        origem = (a.get("pais") or "Internacional") if a.get("internacional") else "Nacional"
        vals = [a["titulo"], a.get("segmento") or "-", a["preco"], a["preco_de"], a["nota"],
                a["vendidos"], "Sim" if a["full"] else "Nao", origem,
                "Sim" if a["loja_oficial"] else "Nao",
                "Sim" if a["frete_gratis"] else "Nao", a["tier"], a["vendedor"], None,
                p["preco"] / 100, p["entrega"] / 100, p["fornecedor"] / 100,
                p["marca"] / 100, a["score"] / 100]
        for col, v in enumerate(vals, 1):
            c = an.cell(i, col, v)
            c.font = Font(name=FT, size=9)
            c.border = Border(bottom=Side("thin", color="D9D9D9"))
            if i % 2 == 0:
                c.fill = PatternFill("solid", fgColor=CINZA)
        for col in (3, 4):
            an.cell(i, col).number_format = '"R$" #,##0.00'
        an.cell(i, 5).number_format = "0.0"
        for col in (14, 15, 16, 17):
            an.cell(i, col).number_format = "0%"
        an.cell(i, 18).number_format = "0.0%"
        an.cell(i, 18).font = Font(name=FT, size=9, bold=True)
        if a.get("internacional"):
            an.cell(i, 8).font = Font(name=FT, size=9, bold=True, color="C00000")
        rotulo = ("abrir anuncio (ad)" if a.get("patrocinado") else "abrir anuncio") \
            if a["url"] else "sem link"
        c = an.cell(i, 13, rotulo)
        if a["url"]:
            c.hyperlink = a["url"]
            c.font = Font(name=FT, size=9, color="0563C1", underline="single")
    an.freeze_panes = "C2"
    if r["elegiveis"]:
        an.auto_filter.ref = "A1:R%d" % (1 + len(r["elegiveis"]))

    # ------------------------------------------------------------ descartados
    ds = wb.create_sheet("Descartados")
    ds.sheet_view.showGridLines = False
    _titulo(ds, "A1", "ANUNCIOS DESCARTADOS PELOS ELIMINATORIOS", "8B4513")
    ds.merge_cells("A1:D1")
    _cabecalho(ds, 3, ["Produto", "Preco", "Origem", "Motivo do descarte",
                       "Vendedor", "Link do anuncio"],
               [52, 13, 14, 42, 18, 16])
    for i, a in enumerate(sorted(r["descartados"], key=lambda x: x["motivo"]), start=4):
        origem = (a.get("pais") or "Internacional") if a.get("internacional") else "Nacional"
        for col, v in enumerate([a["titulo"], a["preco"], origem, a["motivo"],
                                 a["vendedor"], None], 1):
            c = ds.cell(i, col, v)
            c.font = Font(name=FT, size=9)
            c.fill = PatternFill("solid", fgColor=LARANJA)
            c.border = Border(bottom=Side("thin", color="D9D9D9"))
        ds.cell(i, 2).number_format = '"R$" #,##0.00'
        # o link importa aqui: o comprador quer conferir por que o item saiu
        rotulo = ("abrir anuncio (ad)" if a.get("patrocinado") else "abrir anuncio") \
            if a.get("url") else "sem link"
        c = ds.cell(i, 6, rotulo)
        c.fill = PatternFill("solid", fgColor=LARANJA)
        c.border = Border(bottom=Side("thin", color="D9D9D9"))
        if a.get("url"):
            c.hyperlink = a["url"]
            c.font = Font(name=FT, size=9, color="0563C1", underline="single")
        else:
            c.font = Font(name=FT, size=9, italic=True, color="7A8693")
    if r["descartados"]:
        ds.auto_filter.ref = "A3:F%d" % (3 + len(r["descartados"]))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
